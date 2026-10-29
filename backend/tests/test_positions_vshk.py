"""Должности-роли и модуль ВШК: справочник, назначение, ответственные в плане, дашборд, импорт."""
from __future__ import annotations

import io
from datetime import date, timedelta

import pytest

from tests.conftest import auth, pdf_bytes, upload

pytestmark = pytest.mark.asyncio


# --------------------------------------------------------------------- helpers
async def _position_id(client, title: str) -> str:
    rows = (await client.get("/meta/positions")).json()
    for r in rows:
        if r["title"].lower() == title.lower():
            return r["id"]
    raise AssertionError(f"Должность {title!r} не найдена в сиде")


def _vshk_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("ПЛАН ВНУТРИШКОЛЬНОГО КОНТРОЛЯ НА 2026-2027 УЧЕБНЫЙ ГОД")
    table = d.add_table(rows=1, cols=4)
    head = table.rows[0].cells
    for i, text in enumerate(["Цель контроля", "Вид контроля", "Ответственный", "Срок и форма подведения итогов"]):
        head[i].text = text
    rows = [
        ("ОКТЯБРЬ",),
        ("Контроль образовательных результатов учащихся",),
        ("Контроль адаптации учеников 1-х классов.", "Классно-обобщающий",
         "Зам.директора по УВР, педагог-психолог", "Совещание при директоре, справка"),
        ("Проверка организации подготовки обучающихся к ГИА.", "Тематический",
         "Администрация", "Совещание при директоре, справка"),
        ("НОЯБРЬ",),
        ("Контроль уровня профессионального мастерства педагогических работников",),
        ("Контроль деятельности вновь прибывших педагогов.", "Тематический (персональный)",
         "Руководители ШМО, зам.директора по НМР", "Совещание при директоре, справка"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- справочник
async def test_positions_seed_and_crud(client, seeded):
    a = await auth(client, "admin@example.com")
    rows = (await a.get("/meta/positions")).json()
    titles = {r["title"] for r in rows}
    assert "Заместитель директора по учебно-воспитательной работе" in titles
    assert "Классный руководитель" in titles and "Руководитель ШМО" in titles

    created = (await a.post("/meta/positions", json={"title": "Руководитель точки роста"})).json()
    assert created["id"]
    renamed = (await a.patch(f"/meta/positions/{created['id']}", json={"title": "Руководитель «Точки роста»"})).json()
    assert renamed["title"] == "Руководитель «Точки роста»"
    assert (await a.delete(f"/meta/positions/{created['id']}")).status_code == 204

    # учитель не управляет справочником
    t = await auth(client, "teacher1@example.com")
    assert (await t.post("/meta/positions", json={"title": "Хочу должность"})).status_code == 403


async def test_position_assigned_and_delete_guard(client, seeded):
    o = await auth(client, "owner@example.com")
    pos_id = await _position_id(o, "Классный руководитель")
    tid = seeded["teacher1@example.com"]
    resp = await o.patch(f"/users/{tid}", json={"position_ids": [pos_id]})
    assert resp.status_code == 200
    assert [p["id"] for p in resp.json()["positions"]] == [pos_id]
    # занятую должность удалить нельзя
    assert (await o.delete(f"/meta/positions/{pos_id}")).status_code == 409


# ------------------------------------------------- план работы: должности и админ
async def test_plan_resp_position_and_admin_reach_dashboards(client, seeded):
    o = await auth(client, "owner@example.com")
    uvr = await _position_id(o, "Заместитель директора по учебно-воспитательной работе")
    tid = seeded["teacher1@example.com"]
    assert (await o.patch(f"/users/{tid}", json={"position_ids": [uvr]})).status_code == 200

    monday = date.today() - timedelta(days=date.today().weekday())
    content = {
        "duty": {"klass": None, "teacher_id": None},
        "notes": [],
        "month_items": [{"text": "Пункт должности", "place": None, "resp": [f"pos:{uvr}"]}],
        "week_items": [{"text": "Пункт администрации", "place": None, "resp": ["admin"]}],
        "days": [],
    }
    plan = (await o.post("/plans", json={"week_start": monday.isoformat(), "content": content})).json()
    pid = plan["id"]
    assert (await o.post(f"/plans/{pid}/publish")).status_code == 200

    published = (await o.get(f"/plans/{pid}")).json()
    titles = [r["title"] for r in published["content"]["month_items"][0]["resp"]]
    assert titles == ["Заместитель директора по учебно-воспитательной работе"]
    assert published["content"]["week_items"][0]["resp"][0]["kind"] == "admin"

    # держатель должности видит «свой» пункт на дашборде
    t = await auth(client, "teacher1@example.com")
    dash = (await t.get("/dashboard")).json()
    kinds = [e["title"] for e in dash["my_plan_events"]]
    assert "Пункт должности" in kinds
    assert "Пункт администрации" not in kinds

    # администрация видит и должностной пункт? нет — только «admin» и свои
    a = await auth(client, "admin@example.com")
    dash_a = (await a.get("/dashboard")).json()
    titles_a = [e["title"] for e in dash_a["my_plan_events"]]
    assert "Пункт администрации" in titles_a
    assert "Пункт должности" not in titles_a

    # уведомление о публикации дошло до держателя должности
    notif = (await t.get("/notifications")).json()
    assert any("план работы" in n["title"].lower() for n in notif["items"])


async def test_plan_parse_recognises_positions(client, seeded):
    import docx

    a = await auth(client, "admin@example.com")
    d = docx.Document()
    d.add_paragraph("План работы на неделю 26.10.2026г. – 31.10.2026г.")
    d.add_paragraph("В течение месяца: 1. Проверка дневников. Отв.: зам.директора по УВР, Администрация")
    buf = io.BytesIO()
    d.save(buf)
    resp = await a.post(
        "/plans/parse-document",
        files={"file": ("plan.docx", io.BytesIO(buf.getvalue()),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    refs = resp.json()["content"]["month_items"][0]["resp"]
    assert "admin" in refs
    assert any(r.startswith("pos:") for r in refs)


# ------------------------------------------------------------------------ ВШК
async def test_vshk_manual_flow_statuses_and_dashboard(client, seeded):
    o = await auth(client, "owner@example.com")
    uvr = await _position_id(o, "Заместитель директора по учебно-воспитательной работе")
    tid = seeded["teacher1@example.com"]
    assert (await o.patch(f"/users/{tid}", json={"position_ids": [uvr]})).status_code == 200

    year = date.today().year if date.today().month >= 8 else date.today().year - 1
    plan = (await o.post("/vshk", json={"year_start": year})).json()
    pid = plan["id"]
    assert plan["status"] == "DRAFT"

    item = (await o.post(f"/vshk/{pid}/items", json={
        "month": date.today().month,
        "section": "Контроль образовательных результатов",
        "title": "Контроль адаптации учеников 1-х классов.",
        "control_type": "Классно-обобщающий",
        "resp": [f"pos:{uvr}"],
        "summary": "Совещание при директоре, справка",
    })).json()
    assert item["resp_titles"] == ["Заместитель директора по учебно-воспитательной работе"]
    assert item["status"] == "PLANNED"

    # черновик учителю не виден
    t = await auth(client, "teacher1@example.com")
    assert (await t.get(f"/vshk/{pid}")).status_code == 403
    assert (await t.get("/dashboard")).json()["my_vshk"] == []

    assert (await o.post(f"/vshk/{pid}/publish")).status_code == 200
    dash = (await t.get("/dashboard")).json()
    assert [i["title"] for i in dash["my_vshk"]] == ["Контроль адаптации учеников 1-х классов."]

    # статус и результат
    upd = (await o.patch(f"/vshk/items/{item['id']}", json={"status": "IN_WORK"})).json()
    assert upd["status"] == "IN_WORK"
    res = await o.post(
        f"/vshk/items/{item['id']}/result",
        data={"result_text": "Справка по итогам контроля"},
        files=upload("spravka.pdf", pdf_bytes(), "application/pdf"),
    )
    assert res.status_code == 200, res.text
    assert res.json()["result_file_name"] == "spravka.pdf"
    dl = await t.get(f"/vshk/items/{item['id']}/file")
    assert dl.status_code == 200 and dl.headers["content-disposition"].startswith("attachment")

    # удаление мероприятия
    assert (await o.delete(f"/vshk/items/{item['id']}")).status_code == 204
    assert (await o.get(f"/vshk/{pid}")).json()["items"] == []


async def test_vshk_import_docx_and_search(client, seeded):
    o = await auth(client, "owner@example.com")
    year = 2026
    resp = await o.post(
        "/vshk/import",
        data={"year_start": str(year)},
        files={"file": ("vshk.docx", io.BytesIO(_vshk_docx()),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["items_created"] == 3 and out["sections"] == 2
    assert out["warnings"] == []

    plan = (await o.get(f"/vshk/{out['plan_id']}")).json()
    first = plan["items"][0]
    assert first["month"] == 10 and first["section"] == "Контроль образовательных результатов учащихся"
    # «Зам.директора по УВР» и «педагог-психолог» распознаны как должности
    assert sum(1 for r in first["resp"] if r.startswith("pos:")) == 2
    admin_item = plan["items"][1]
    assert "admin" in admin_item["resp"]

    # повторный импорт в черновик заменяет мероприятия
    resp2 = await o.post(
        "/vshk/import",
        data={"year_start": str(year)},
        files={"file": ("vshk.docx", io.BytesIO(_vshk_docx()),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp2.status_code == 200
    assert (await o.get(f"/vshk/{out['plan_id']}")).json()["items"].__len__() == 3

    # публикация → поиск находит мероприятие и у учителя
    assert (await o.post(f"/vshk/{out['plan_id']}/publish")).status_code == 200
    t = await auth(client, "teacher1@example.com")
    found = (await t.get("/search", params={"q": "адаптации"})).json()
    assert found["vshk"], found
    assert found["vshk"][0]["url"].startswith("/vshk?item=")


async def test_vshk_duplicate_year_and_teacher_cannot_manage(client, seeded):
    o = await auth(client, "owner@example.com")
    year = 2031
    assert (await o.post("/vshk", json={"year_start": year})).status_code == 201
    dup = await o.post("/vshk", json={"year_start": year})
    assert dup.status_code == 409
    t = await auth(client, "teacher1@example.com")
    assert (await t.post("/vshk", json={"year_start": 2032})).status_code == 403
